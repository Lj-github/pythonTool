__author__ = 'songbin'
'''创建对应资源翻译的xls 收集对应语言的翻译'''
import os
import xlrd
import xlwt
import chardet
import docx
# resdirList  = [
#     [
#         '/Users/songbin/sanguo/art_pikachu/翻译美术（越南）/越南版翻译',
#         '/Users/songbin/sanguo/art_pikachu/翻译美术图',
#         '/Users/songbin/sanguo/art_pikachu/程天游',
#         '/Users/songbin/sanguo/art_pikachu/绿洲',
#         '/Users/songbin/Downloads/Language/18年1月至4月翻译'
#     ],
#     ['/Users/songbin/sanguo/art_pikachu/翻译美术图（俄文版）'],
#     [
#         '/Users/songbin/sanguo/art_pikachu/翻译美术（越南）/越语版',
#         '/Users/songbin/sanguo/art_pikachu/翻译美术（越南）/越南_0423'
#     ],
#     ['/Users/songbin/sanguo/art_pikachu/翻译美术（法国）'],
#     ['/Users/songbin/sanguo/art_pikachu/翻译美术图（德国）'],
# ]
resdirList = [
    [
        '/Users/songbin/Downloads/Language/18年1月至4月翻译',
    ],
    [],
    [],
    [],
    ['/Users/songbin/sanguo/art_pikachu/翻译美术图（德国）'],
]
translateList = [
    1,
    2,
    3,
    4,
    5
]
titleList = ['fileName','English','Russion','Vietnam','French','Germany','TraditionalChinese']

imgTranslateXlsPth='/Users/songbin/Downloads/Language/translateImgData.xls'

keyNameType = ['png','jpg']
valueType = ['txt','docx']


def makeMap(srcDir):
    print('makeMap:',srcDir)
    print('translateId:',translateId)
    for root, dirs, files in os.walk(srcDir):
            for OneFileName in files:
                # print(os.path.join(root,OneFileName))
                fileType = OneFileName.split('.')[-1:][0]
                if fileType in keyNameType:
                    value = imgMap.get(OneFileName)
                    if value:
                        continue
                    else:
                        nameNoneType = OneFileName.split('.')[:1][0]
                        value1 = imgMap.get(nameNoneType)
                        if value1:
                            imgMap[OneFileName]=value1
                            imgMap.pop(nameNoneType)
                            # print('nameNoneType:',nameNoneType)
                        else:
                            imgMap[OneFileName]={}

                elif fileType in valueType:
                    lines=''
                    filePth=os.path.join(root,OneFileName)
                    if (fileType == valueType[0]):
                        f=open(filePth,'rb')
                        txtEncodeInfo = chardet.detect(f.read())
                        encodeStr=txtEncodeInfo.get('encoding')
                        if not encodeStr:
                            encodeStr = 'latin1'
                        # print('chardet:',encodeStr)

                        with open(filePth, 'r',32,encodeStr) as r:
                            lines=r.read()

                    elif (fileType == valueType[1]):
                        doc = docx.Document(filePth)
                        for para in doc.paragraphs:
                            if para.text and para.text != '':
                                str = para.text
                                if (lines == ''):
                                    lines=str
                                else:
                                    lines=lines + ';' +str
                                print('docx:',para.text)


                    if (lines == ''):
                        continue
                    nameNoneType = OneFileName.split('.')[:1][0]
                    nameNoneType = nameNoneType.split('_rus')[:1][0]

                    if (imgMap.get(nameNoneType + '.png') or imgMap.get(nameNoneType + '.png') == {}):
                        imgMap[nameNoneType + '.png'][translateId] = lines
                    elif (imgMap.get(nameNoneType + '.jpg') or imgMap.get(nameNoneType + '.jpg') == {}):
                        imgMap[nameNoneType + '.jpg'][translateId] = lines
                    else:
                        if not imgMap.get(nameNoneType):
                            imgMap[nameNoneType] = {}
                        imgMap[nameNoneType][translateId] = lines


def writeExcel():
    if (os.path.exists(imgTranslateXlsPth)):
        table_translate = xlrd.open_workbook(imgTranslateXlsPth)
        table = table_translate.sheet_by_index(0)
        """获取table工作表总行数"""
        nrows = table.nrows
        '''获取table工作表总列数'''
        ncols = table.ncols
    else:
        nrows = 0

    """创建一个excel文件"""
    workbook = xlwt.Workbook()
    newSheet = workbook.add_sheet('sheet0',cell_overwrite_ok=True)
    # print("imgMap:",imgMap)
    col0 = 0
    for colName in titleList:
        newSheet.write(0, col0, colName)
        col0 = col0 + 1
    findValues=[]
    if (nrows == 0):
        nrows = 1
    else:
        for i in range(nrows):
            if i == 0:
                continue
            fileName = table.cell_value(i, 0)
            value = imgMap.get(fileName)
            for j in range(ncols):
                if j == 0:
                    newSheet.write(i, j, table.cell_value(i, j))
                    continue

                if (value):
                    findValues.append(fileName)

                    writeColList = []
                    for transId,transStr in value.items():
                        newSheet.write(i, transId, transStr)
                        writeColList.append(transId)
                    for _col in range(5):
                        if _col == 0:
                            continue
                        if _col in writeColList:
                            continue
                        newSheet.write(i, _col, table.cell_value(i, _col))
                    break

                else:
                    cell_value = table.cell_value(i, j)
                    newSheet.write(i, j, cell_value)

    for _fileName, _fileValues in imgMap.items():
        if _fileName in findValues:
            continue
        isNewLineNull = True
        if _fileValues == {}:
            isNewLineNull = False
        for transId,transStr in _fileValues.items():
            if transStr != '' and transStr:
                isNewLineNull = False
                print(transStr)

        if not isNewLineNull:
            continue

        newSheet.write(nrows, 0, _fileName)
        for transId,transStr in _fileValues.items():
            newSheet.write(nrows, transId, transStr)
        nrows = nrows + 1

    fpath,fname=os.path.split(imgTranslateXlsPth)
    if not os.path.exists(fpath):
        '''创建路径'''
        os.makedirs(fpath)
    print('保存文件 ',imgTranslateXlsPth)

    workbook.save(imgTranslateXlsPth)


if __name__ == "__main__":
    try:
        global imgMap
        imgMap = {}
        global translateId
        resDirCount = len(resdirList)
        for k in range(resDirCount):
            dirList = resdirList[k]
            translateId = translateList[k]
            for dir in dirList:
                makeMap(dir)
        writeExcel()

    except Exception as e:
        print("异常" + e)